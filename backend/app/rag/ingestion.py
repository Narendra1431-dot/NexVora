import io
import zipfile
from fastapi import HTTPException
from pypdf import PdfReader

MIMES={'pdf':{'application/pdf','application/octet-stream'},'txt':{'text/plain','application/octet-stream'},
       'docx':{'application/vnd.openxmlformats-officedocument.wordprocessingml.document','application/octet-stream'}}

def validate_file(name,mime,data):
    extension=name.lower().rsplit('.',1)[-1]
    if extension not in MIMES or mime not in MIMES[extension]:raise HTTPException(415,'Upload PDF, UTF-8 TXT or DOCX.')
    if not data or len(data)>8*1024*1024:raise HTTPException(413,'Upload a non-empty document up to 8 MB.')
    if extension=='pdf' and not data.startswith(b'%PDF-'):raise HTTPException(422,'Invalid PDF signature.')
    if extension=='docx':
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if 'word/document.xml' not in archive.namelist() or sum(f.file_size for f in archive.infolist())>32*1024*1024:raise ValueError()
        except Exception:raise HTTPException(422,'Invalid or oversized DOCX archive.')
    if extension=='txt':
        try:data.decode('utf-8-sig')
        except UnicodeError:raise HTTPException(422,'TXT must be UTF-8.')
    return extension

def parse(data,extension):
    pages=[]
    if extension=='pdf':
        reader=PdfReader(io.BytesIO(data))
        if reader.is_encrypted:raise ValueError('Encrypted PDFs are not supported.')
        if len(reader.pages)>600:raise ValueError('Document exceeds 600 pages.')
        for i,page in enumerate(reader.pages):
            text=page.extract_text() or ''
            if len(text.strip())<20:
                # OCR is opt-in and only attempted on pages without text.
                import os
                if os.getenv('ENABLE_DOCUMENT_OCR','false').lower()=='true':
                    import fitz,pytesseract
                    from PIL import Image
                    with fitz.open(stream=data,filetype='pdf') as pdf:
                        pix=pdf[i].get_pixmap(matrix=fitz.Matrix(1.5,1.5))
                        text=pytesseract.image_to_string(Image.frombytes('RGB',[pix.width,pix.height],pix.samples))
                if len(text.strip())<20:raise ValueError('This PDF needs OCR. Enable the optional OCR worker or upload a text-based PDF.')
            pages.append({'page':i+1,'text':text})
    elif extension=='docx':
        from docx import Document
        doc=Document(io.BytesIO(data))
        text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
        # DOCX pagination is renderer-dependent; page=0 explicitly means unknown.
        pages=[{'page':0,'text':text}]
    else:pages=[{'page':1,'text':data.decode('utf-8-sig')}]
    if not any(p['text'].strip() for p in pages):raise ValueError('No extractable text.')
    if sum(len(p['text']) for p in pages)>2_000_000:raise ValueError('Extracted text exceeds the indexing limit.')
    return pages
